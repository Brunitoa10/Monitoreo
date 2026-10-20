"""Reporte Excel con formato: resumen, recomendaciones, detalle y gráficos nativos."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from nemesis.reporting.base import ReportDataset, ReportExporter
from nemesis.reporting.formatting import LEVEL_LABEL, kpi_rows, status_label

HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="0F172A")
MUTED_FONT = Font(size=9, color="64748B")
THIN = Border(bottom=Side(style="thin", color="E2E8F0"))
GOOD = PatternFill("solid", fgColor="DCFCE7")
WARN = PatternFill("solid", fgColor="FEF3C7")
BAD = PatternFill("solid", fgColor="FEE2E2")
MAX_RAW_ROWS = 200_000

COLUMN_LABELS = {
    "site": "Sitio", "type": "Tipo", "equipos": "Equipos", "mediciones": "Mediciones",
    "disponibilidad_pct": "Disponibilidad %", "latencia_prom_ms": "Latencia prom (ms)",
    "latencia_p95_ms": "Latencia p95 (ms)", "latencia_max_ms": "Latencia máx (ms)",
    "perdida_prom_pct": "Pérdida prom %", "jitter_prom_ms": "Jitter prom (ms)", "incidentes": "Incidentes",
    "downtime_min": "Downtime (min)", "mttr_min": "MTTR (min)", "equipo": "Equipo", "tipo": "Tipo",
    "sitio": "Sitio", "ip": "IP", "estado_actual": "Estado actual", "saltos": "Saltos",
    "device_name": "Equipo", "status": "Severidad", "started_at": "Inicio", "ended_at": "Fin",
    "duration_min": "Duración (min)", "is_open": "Abierto", "message": "Detalle",
    "saltos_estimados_ttl": "Saltos estimados (TTL)", "saltos_traceroute": "Saltos traceroute",
    "rutas_distintas": "Rutas distintas", "llega_destino": "Llega a destino", "rtt_ruta_ms": "RTT ruta (ms)",
    "ultimo_traceroute": "Último traceroute", "fecha_hora": "Fecha/hora", "disponibles": "Mediciones OK", "ts": "Fecha/hora", "online": "Responde", "latency_ms": "Latencia (ms)",
    "jitter_ms": "Jitter (ms)", "loss_pct": "Pérdida %", "hops": "Saltos", "device_id": "ID equipo",
}


def _write_df(ws: Worksheet, df: pd.DataFrame, start_row: int = 1, table_name: str | None = None) -> int:
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_bool_dtype(df[col]) or df[col].dtype == object:
            df[col] = df[col].map(lambda v: "Sí" if v is True else "No" if v is False else v)
    headers = [COLUMN_LABELS.get(c, c) for c in df.columns]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=j, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            if isinstance(value, float) and value != value:
                value = None
            elif isinstance(value, pd.Timestamp):
                value = None if pd.isna(value) else value.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = THIN
            if hasattr(value, "year"):
                cell.number_format = "dd/mm/yyyy hh:mm"
    end_row = start_row + len(df)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for j, h in enumerate(headers, start=1):
        sample = [len(str(h))] + [len(str(ws.cell(row=r, column=j).value or "")) for r in range(start_row + 1, min(end_row, start_row + 200) + 1)]
        ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, max(sample) + 2))
    return end_row


def _col(df: pd.DataFrame, name: str) -> str | None:
    if name not in df.columns:
        return None
    return get_column_letter(list(df.columns).index(name) + 1)


def _availability_format(ws: Worksheet, df: pd.DataFrame, start: int, end: int) -> None:
    letter = _col(df, "disponibilidad_pct")
    if not letter or end <= start:
        return
    rng = f"{letter}{start + 1}:{letter}{end}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F87171",
                                                      end_type="max", end_color="FFFFFF"))
    for c in ("latencia_p95_ms", "perdida_prom_pct", "downtime_min"):
        lc = _col(df, c)
        if lc:
            ws.conditional_formatting.add(f"{lc}{start + 1}:{lc}{end}", ColorScaleRule(
                start_type="min", start_color="FFFFFF", end_type="max", end_color="F87171"))


def _bar_chart(ws: Worksheet, df: pd.DataFrame, cat_col: str, title: str, anchor: str, start: int, end: int) -> None:
    if end <= start:
        return
    chart = BarChart()
    chart.type, chart.style, chart.title = "bar", 10, title
    chart.y_axis.title = "Disponibilidad %"
    chart.y_axis.scaling.max = 100
    data = Reference(ws, min_col=list(df.columns).index("disponibilidad_pct") + 1, min_row=start, max_row=end)
    cats = Reference(ws, min_col=list(df.columns).index(cat_col) + 1, min_row=start + 1, max_row=end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.height, chart.width = max(7, 0.6 * (end - start) + 3), 16
    ws.add_chart(chart, anchor)


class ExcelExporter(ReportExporter):
    key = "excel"
    label = "Excel (detalle + gráficos)"
    extension = "xlsx"
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def export(self, data: ReportDataset) -> bytes:
        wb = Workbook()
        self._summary(wb.active, data)

        ws = wb.create_sheet("Recomendaciones")
        rec = pd.DataFrame([{"Nivel": LEVEL_LABEL[i.level], "Hallazgo": i.title, "Detalle": i.detail,
                             "Recomendación": i.recommendation} for i in data.insights])
        if not rec.empty:
            _write_df(ws, rec)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        for title, df, cat in (("Por sitio", data.by_site, "site"), ("Por tipo", data.by_type, "type")):
            ws = wb.create_sheet(title)
            end = _write_df(ws, df)
            _availability_format(ws, df, 1, end)
            if not df.empty:
                _bar_chart(ws, df, cat, f"Disponibilidad {title.lower()}", f"{get_column_letter(len(df.columns) + 2)}2", 1, end)

        ws = wb.create_sheet("Equipos")
        devices = data.devices.drop(columns=["device_id"], errors="ignore").copy()
        if "estado_actual" in devices:
            devices["estado_actual"] = devices["estado_actual"].map(status_label)
        end = _write_df(ws, devices)
        _availability_format(ws, devices, 1, end)

        self._audit_sheets(wb, data)

        ws = wb.create_sheet("Saltos")
        _write_df(ws, data.hops.drop(columns=["device_id"], errors="ignore"))

        ws = wb.create_sheet("Tendencia")
        series = data.availability_series.merge(
            data.latency_series.groupby("ts", as_index=False)["latencia_prom_ms"].mean()
            if not data.latency_series.empty else pd.DataFrame(columns=["ts", "latencia_prom_ms"]),
            on="ts", how="left")
        end = _write_df(ws, series)
        if end > 2:
            line = LineChart()
            line.title, line.height, line.width = "Disponibilidad en el tiempo", 8, 22
            line.add_data(Reference(ws, min_col=2, min_row=1, max_row=end), titles_from_data=True)
            line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=end))
            line.y_axis.title = "%"
            ws.add_chart(line, "E2")

        if not data.hourly.empty:
            ws = wb.create_sheet("Por hora")
            hourly = data.hourly.drop(columns=["device_id"], errors="ignore").tail(MAX_RAW_ROWS)
            end = _write_df(ws, hourly)
            _availability_format(ws, hourly, 1, end)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @staticmethod
    def _audit_sheets(wb: Workbook, data: ReportDataset) -> None:
        """Auditoría completa: una fila por falla con fecha/hora exactas, y totales por equipo."""
        ws = wb.create_sheet("Auditoría de fallas", 1)
        ws["A1"] = "Auditoría de fallas"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = ("Caída = primer chequeo sin respuesta · Recuperación = primer chequeo OK posterior · "
                    "Último OK = último chequeo correcto antes de la caída (el corte ocurrió entre ambos).")
        ws["A2"].font = MUTED_FONT
        end = _write_df(ws, data.audit, start_row=4)
        if end > 4:
            col = list(data.audit.columns).index("Estado") + 1
            letter = get_column_letter(col)
            ws.conditional_formatting.add(f"{letter}5:{letter}{end}", CellIsRule(
                operator="equal", formula=['"Sin recuperar"'], fill=BAD))
            ws.conditional_formatting.add(f"{letter}5:{letter}{end}", CellIsRule(
                operator="equal", formula=['"Recuperado"'], fill=GOOD))
        ws2 = wb.create_sheet("Fallas por equipo", 2)
        _write_df(ws2, data.audit_summary)

    @staticmethod
    def _summary(ws: Worksheet, data: ReportDataset) -> None:
        ws.title = "Resumen"
        ws["A1"] = data.title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"{data.organization} · Período: {data.filter.label} · Generado: {data.generated_at:%d/%m/%Y %H:%M}"
        ws["A2"].font = MUTED_FONT
        if data.filter.sites or data.filter.types:
            ws["A3"] = f"Sitios: {', '.join(data.filter.sites) or 'todos'} · Tipos: {', '.join(data.filter.types) or 'todos'}"
            ws["A3"].font = MUTED_FONT
        for j, h in enumerate(("Indicador", "Valor"), start=1):
            c = ws.cell(row=5, column=j, value=h)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
        for i, (label, value) in enumerate(kpi_rows(data.kpis), start=6):
            ws.cell(row=i, column=1, value=label).border = THIN
            cell = ws.cell(row=i, column=2, value=value)
            cell.border, cell.font = THIN, Font(bold=True)
            cell.alignment = Alignment(horizontal="right")
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 22
