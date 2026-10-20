"""Auditoría de fallas: cuándo cayó cada equipo (fecha y hora con segundos), cuándo volvió
y cuánto tiempo estuvo con falla. Exportable a Excel y CSV."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook

from nemesis.container import Container
from nemesis.domain.models import format_duration
from nemesis.reporting.excel_exporter import _write_df
from nemesis.services.metrics import ReportFilter
from nemesis.ui import components as ui

AUDIT_COLUMNS = {
    "N°": st.column_config.NumberColumn("N°", format="%d", width="small"),
    "Equipo": st.column_config.TextColumn("Equipo", width="medium"),
    "Último OK": st.column_config.TextColumn("Último OK", help="Último chequeo correcto antes de la caída"),
    "Fecha caída": st.column_config.TextColumn("Fecha caída"),
    "Hora caída": st.column_config.TextColumn("Hora caída", help="Primer chequeo sin respuesta (hh:mm:ss)"),
    "Fecha recuperación": st.column_config.TextColumn("Fecha recuperación"),
    "Hora recuperación": st.column_config.TextColumn("Hora recuperación", help="Primer chequeo OK posterior"),
    "Duración": st.column_config.TextColumn("Duración", width="medium"),
    "Duración (s)": None,
    "Detalle": st.column_config.TextColumn("Detalle", width="large"),
}


def _excel(audit: pd.DataFrame, summary: pd.DataFrame, f: ReportFilter, org: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría de fallas"
    ws["A1"] = f"Auditoría de fallas · {org}"
    ws["A2"] = f"Período: {f.label} · Generado: {datetime.now():%d/%m/%Y %H:%M:%S}"
    _write_df(ws, audit, start_row=4)
    _write_df(wb.create_sheet("Fallas por equipo"), summary)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(c: Container, f: ReportFilter) -> None:
    ui.header("Auditoría de fallas", f"Registro de caídas y recuperaciones · {f.label}")
    audit = c.metrics.audit_log(f)
    summary = c.metrics.audit_summary(f)
    interval = c.settings.collector.interval_seconds
    st.caption(f"Precisión: chequeo cada {interval} s. La caída ocurrió entre «Último OK» y «Hora caída»; "
               "la recuperación, en la «Hora recuperación» (primer chequeo correcto).")

    if audit.empty:
        ui.empty("Sin fallas registradas en el período. 👌")
        return

    open_n = int((audit["Estado"] == "Sin recuperar").sum())
    cols = st.columns(5)
    with cols[0]:
        ui.kpi("Fallas", str(len(audit)), f"{(audit['Severidad'] == 'Caída').sum()} caídas")
    with cols[1]:
        ui.kpi("Sin recuperar", str(open_n), "en curso ahora", "bad" if open_n else "ok")
    with cols[2]:
        ui.kpi("Tiempo total con falla", format_duration(audit["Duración (s)"].sum()), "suma de todas las fallas")
    with cols[3]:
        ui.kpi("Falla más larga", format_duration(audit["Duración (s)"].max()),
               str(audit.loc[audit["Duración (s)"].idxmax(), "Equipo"]))
    with cols[4]:
        ui.kpi("Equipos afectados", str(audit["Equipo"].nunique()), "")

    st.write("")
    a, b, d = st.columns([1, 1, 2])
    estado = a.selectbox("Estado", ["Todos", "Sin recuperar", "Recuperado"])
    severidad = b.selectbox("Tipo de falla", ["Todas", "Caída", "Degradación"])
    search = d.text_input("Buscar", placeholder="equipo, sitio o detalle")
    view = audit
    if estado != "Todos":
        view = view[view["Estado"] == estado]
    if severidad != "Todas":
        view = view[view["Severidad"] == severidad]
    if search:
        mask = pd.Series(False, index=view.index)
        for col in ("Equipo", "Sitio", "Tipo", "Detalle"):
            mask |= view[col].astype(str).str.contains(search, case=False, na=False)
        view = view[mask]

    shown = view.assign(Estado=view["Estado"].map({"Sin recuperar": "🔴 Sin recuperar", "Recuperado": "🟢 Recuperado"}))
    st.dataframe(shown, column_config=AUDIT_COLUMNS, hide_index=True, width="stretch",
                 height=min(560, 36 * len(view) + 40))

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    e1, e2 = st.columns(2)
    e1.download_button("⬇️ Exportar a Excel", _excel(view, summary, f, c.settings.app.organization),
                       file_name=f"auditoria_fallas_{stamp}.xlsx", width="stretch",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    e2.download_button("⬇️ Exportar a CSV", view.to_csv(index=False, sep=";").encode("utf-8-sig"),
                       file_name=f"auditoria_fallas_{stamp}.csv", mime="text/csv", width="stretch")

    st.subheader("Fallas por equipo")
    st.dataframe(summary.drop(columns=["Segundos"]), hide_index=True, width="stretch",
                 height=min(420, 36 * len(summary) + 40))
